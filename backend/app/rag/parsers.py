"""Parser：把不同格式的文件统一转成 ParsedDocument（若干带页码的文本段）。

新增格式（PPTX、HTML、网页抓取、飞书文档…）只需要注册一个新的 parser 函数。
"""

from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ParsedSection:
    text: str
    page: int | None = None      # PDF 页码（从 1 开始）；其他格式为 None


@dataclass
class ParsedDocument:
    sections: list[ParsedSection] = field(default_factory=list)
    page_count: int | None = None
    is_markdown: bool = False     # 文本中是否保留了 # 标题结构，供 Markdown 切片器使用


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):  # 兼容 Windows 下常见的 GBK 编码 TXT
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def parse_text(path: Path) -> ParsedDocument:
    return ParsedDocument([ParsedSection(_read_text(path))])


def parse_markdown(path: Path) -> ParsedDocument:
    return ParsedDocument([ParsedSection(_read_text(path))], is_markdown=True)


def parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    sections = [ParsedSection(page.extract_text() or "", page=i + 1)
                for i, page in enumerate(reader.pages)]
    if not any(s.text.strip() for s in sections):
        raise ValueError("PDF 中没有可提取的文字（可能是扫描件，需要 OCR，见 Roadmap）")
    return ParsedDocument(sections, page_count=len(reader.pages))


def parse_docx(path: Path) -> ParsedDocument:
    """把 Word 的标题样式转成 Markdown 标题，这样后续可以按章节切片、在引用里显示章节名。"""
    import docx
    d = docx.Document(str(path))
    lines: list[str] = []
    list_no = 0  # 编号列表的序号在 XML 里是自动生成的，p.text 拿不到，这里补回来
    body = d.element.body
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(el, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") or style.startswith("标题"):
                level = int("".join(ch for ch in style if ch.isdigit()) or "1")
                lines.append(f"{'#' * min(level + 1, 6)} {text}")  # Title 占用一级标题
            elif style == "title":
                lines.append(f"# {text}")
            elif style.startswith("list number"):
                list_no += 1
                lines.append(f"{list_no}. {text}")
                continue
            elif style.startswith("list"):
                lines.append(f"- {text}")
            else:
                lines.append(text)
            list_no = 0
        elif tag == "tbl":
            t = docx.table.Table(el, d)
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(" | ".join(cells))
    return ParsedDocument([ParsedSection("\n\n".join(lines))], is_markdown=True)


PARSERS: dict[str, Callable[[Path], ParsedDocument]] = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "md": parse_markdown,
    "markdown": parse_markdown,
    "txt": parse_text,
}

SUPPORTED_TYPES = sorted(PARSERS)


def detect_type(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def parse_file(path: Path, file_type: str) -> ParsedDocument:
    parser = PARSERS.get(file_type)
    if parser is None:
        raise ValueError(f"暂不支持 .{file_type} 文件，支持：{', '.join(SUPPORTED_TYPES)}")
    return parser(path)
